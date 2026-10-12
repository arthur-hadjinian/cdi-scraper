import hashlib,json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as E
import openpyxl

job=Path('outputs/linkedin-pilot')
d=json.loads((job/'export.json').read_text(encoding='utf-8'))
assert hashlib.sha256(Path(d['input']).read_bytes()).hexdigest()==d['sha256'], 'Original workbook changed'
w=openpyxl.load_workbook(job/'approved-linkedin-pilot.xlsx',read_only=False,data_only=False)
s=w['Passed - LinkedIn']
assert s.max_row==len(d['rows'])+6
assert s.max_column==41
headers=[s.cell(6,i).value for i in range(1,42)]
for i,row in enumerate(s.iter_rows(min_row=7,values_only=True)):
    expected=d['rows'][i]
    actual=dict(zip(headers,row))
    for h in d['headers']:
        assert actual[h]==expected[h],(i,h,actual[h],expected[h])
    m=d['results'].get(expected['License_Number'])
    if m:
        assert actual['LinkedIn match']==m['status']
        assert (actual['LinkedIn URL'] or '')==m['linkedin_url']
        assert (actual['Possible LinkedIn URL'] or '')==m['candidate_url']
        assert actual['Match evidence']==m['evidence']
        assert actual['Recorded searches']==len(m['queries'])
        assert actual['LinkedIn checked'].date().isoformat()==m['checked_on']
    else: assert actual['LinkedIn match']=='Not researched'
assert s.freeze_panes=='B7'
counts=Counter(m['status'] for m in d['results'].values())
assert counts=={'Strong match':7,'Possible match':5,'Not found':38}
errors={'#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!','#SPILL!'}
assert not any(c.value in errors for row in s for c in row if c.data_type=='e')
with ZipFile(job/'approved-linkedin-pilot.xlsx') as z:
    root=E.fromstring(z.read('xl/worksheets/sheet1.xml'))
    links=root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}hyperlinks')
    assert len(links)==5244
print(json.dumps({'original_unchanged':True,'all_original_values_preserved':len(d['rows'])*len(d['headers']),'first_50':dict(counts),'native_hyperlinks':len(links),'formula_errors':0},indent=2))
